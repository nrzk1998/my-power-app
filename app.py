# -*- coding: utf-8 -*-
import io
import os

from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st

import analyzer
import plotter
import scraper


st.set_page_config(page_title="電力分析レポート", layout="wide")


PASSWORD = "acorn"
STATE_DEFAULTS = {
    'analyzed': False,
    'final_df': None,
    'fig_report': None,
    'fig_report_buf': None,
    'fig_cal_buf': None,
    'input_filename': "",
    'k_final': None,
    'target_station': "",
}


def inject_styles():
    st.markdown(
        """
        <style>
        .stApp {
            background-color: #f5f5f2;
        }
        .block-container {
            padding-top: 2.5rem;
            padding-bottom: 3.5rem;
            max-width: 1200px;
        }
        .hero-panel {
            padding: 2.5rem 2.5rem;
            border-radius: 3px;
            color: #1a1a1a;
            background: #ffffff;
            border: 1px solid #e0dedd;
            margin-bottom: 1.5rem;
        }
        .hero-kicker {
            font-size: 0.72rem;
            letter-spacing: 0.18em;
            text-transform: uppercase;
            color: #4f7a5e;
            margin-bottom: 0.75rem;
            font-weight: 600;
        }
        .hero-title {
            font-size: 2rem;
            line-height: 1.2;
            font-weight: 600;
            margin: 0;
            color: #1a1a1a;
        }
        .hero-copy {
            margin-top: 1rem;
            font-size: 0.93rem;
            line-height: 1.85;
            color: #555550;
            max-width: 720px;
        }
        .status-grid {
            display: grid;
            grid-template-columns: repeat(4, minmax(0, 1fr));
            gap: 1rem;
            margin: 1.25rem 0 1.5rem;
        }
        .status-card {
            background: #ffffff;
            border: 1px solid #e0dedd;
            border-radius: 3px;
            padding: 1.25rem 1.5rem;
        }
        .status-label {
            font-size: 0.7rem;
            text-transform: uppercase;
            letter-spacing: 0.14em;
            color: #9a9992;
            margin-bottom: 0.5rem;
            font-weight: 600;
        }
        .status-value {
            font-size: 1.35rem;
            color: #1a1a1a;
            font-weight: 600;
        }
        .section-card {
            background: #ffffff;
            border: 1px solid #e0dedd;
            border-radius: 3px;
            padding: 1.5rem;
        }
        .empty-state {
            padding: 2.5rem;
            border-radius: 3px;
            background: #ffffff;
            border: 1px solid #e0dedd;
        }
        .login-panel {
            max-width: 480px;
            margin: 4rem auto 0;
            padding: 2.5rem;
            border-radius: 3px;
            background: #ffffff;
            border: 1px solid #e0dedd;
        }
        [data-testid="stSidebar"] {
            background: #ffffff;
            border-right: 1px solid #e0dedd;
        }
        [data-testid="stSidebar"] * {
            color: #1a1a1a;
        }
        [data-testid="stSidebar"] .stButton button {
            background: #4f7a5e;
            color: #ffffff;
            border: none;
            border-radius: 2px;
            font-weight: 600;
            letter-spacing: 0.03em;
        }
        [data-testid="stSidebar"] .stButton button:hover {
            background: #3e6349 !important;
        }
        [data-testid="stSidebar"] .stDownloadButton button {
            background: #4f7a5e !important;
            color: #ffffff !important;
            -webkit-text-fill-color: #ffffff !important;
            border: none !important;
            border-radius: 2px;
            font-weight: 600;
            letter-spacing: 0.03em;
        }
        [data-testid="stSidebar"] .stDownloadButton button:hover {
            background: #3e6349 !important;
            color: #ffffff !important;
            -webkit-text-fill-color: #ffffff !important;
        }
        [data-testid="stSidebar"] .stDownloadButton button p,
        [data-testid="stSidebar"] .stDownloadButton button span,
        [data-testid="stSidebar"] .stDownloadButton button div {
            color: #ffffff !important;
            -webkit-text-fill-color: #ffffff !important;
            font-weight: 600 !important;
        }
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] button,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] small,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] span,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] div,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] small,
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] span,
        [data-testid="stSidebar"] [data-testid="stFileUploaderFileName"] {
            color: #1a1a1a !important;
            -webkit-text-fill-color: #1a1a1a !important;
            opacity: 1 !important;
        }
        [data-testid="stSidebar"] div[data-baseweb="input"] input,
        [data-testid="stSidebar"] div[data-baseweb="base-input"] input,
        [data-testid="stSidebar"] .stNumberInput input {
            color: #1a1a1a !important;
            -webkit-text-fill-color: #1a1a1a !important;
            background: #f5f5f2 !important;
            opacity: 1 !important;
        }
        [data-testid="stSidebar"] input::placeholder {
            color: #777770 !important;
            -webkit-text-fill-color: #777770 !important;
            opacity: 1 !important;
        }
        div[data-baseweb="input"] input,
        div[data-baseweb="select"] input,
        div[data-baseweb="base-input"] input {
            border-radius: 2px;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def ensure_session_state():
    for key, value in STATE_DEFAULTS.items():
        if key not in st.session_state:
            st.session_state[key] = value


def password_entered():
    if st.session_state.get('password') == PASSWORD:
        st.session_state['password_correct'] = True
        del st.session_state['password']
    else:
        st.session_state['password_correct'] = False


def check_password():
    if st.session_state.get('password_correct'):
        return True

    st.markdown(
        """
        <div class="login-panel">
            <div class="hero-kicker">Power Analysis</div>
            <h1 class="hero-title">電力分析ダッシュボード</h1>
            <p class="hero-copy">
                解析機能に入る前にパスワードを入力してください。
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    prompt = "パスワードを入力してください"
    if 'password_correct' in st.session_state and not st.session_state['password_correct']:
        prompt = "パスワードが違います。再入力してください"

    st.text_input(prompt, type="password", on_change=password_entered, key="password")
    if 'password_correct' in st.session_state and not st.session_state['password_correct']:
        st.error("パスワードが正しくありません")
    return False


def read_power_data(uploaded_file, floor_area):
    df_raw = pd.read_csv(uploaded_file, index_col=0)
    df_raw = df_raw[df_raw.index.notna()]
    df_raw.index = pd.to_datetime(df_raw.index)
    df_raw.index.name = 'Date'

    df_unit = (df_raw * 1000) / floor_area
    return df_raw, df_unit


def create_figure_buffer(figure):
    buf = io.BytesIO()
    figure.savefig(buf, format="png", bbox_inches='tight', dpi=150)
    return buf.getvalue()


def create_calendar_buffer(final_df):
    fig_cal = plotter.create_calendar_report(final_df)
    fig_cal_buf = create_figure_buffer(fig_cal)
    plt.close(fig_cal)
    return fig_cal_buf


def create_word_report_buffer(input_filename, target_station, k_final, report_buf, calendar_buf):
    document = Document()
    document.add_heading('電力消費クラスタリングレポート', level=0)
    document.add_paragraph(f'対象ファイル: {input_filename}')
    document.add_paragraph(f'対象地点: {target_station}')
    document.add_paragraph(f'クラスタ数: {k_final}')

    document.add_heading('クラスタリング結果', level=1)
    document.add_picture(io.BytesIO(report_buf), width=Inches(6.5))

    document.add_heading('カレンダー画像', level=1)
    document.add_picture(io.BytesIO(calendar_buf), width=Inches(6.5))

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def build_analysis_result(uploaded_file, floor_area, target_station, k_input):
    df_raw, df_unit = read_power_data(uploaded_file, floor_area)
    start_date = df_raw.index[0].date()
    end_date = df_raw.index[-1].date()

    df_weather = scraper.get_weather_data(target_station, start_date, end_date)
    if df_weather is None:
        return None

    clusters, k_final = analyzer.perform_clustering(df_unit, k_manual=k_input)
    df_unit = df_unit.copy()
    df_unit['Cluster'] = clusters
    df_raw = df_raw.copy()
    df_raw['Cluster'] = clusters

    report_df = df_unit.join(df_weather, how='inner')
    report_df['DayType'] = report_df['IsHoliday'].map({False: 'Weekday', True: 'Weekend'})

    final_df = df_raw.join(df_weather, how='inner')
    final_df['DayType'] = final_df['IsHoliday'].map({False: 'Weekday', True: 'Weekend'})

    fig_report = plotter.create_combined_report(report_df, k_final)

    return {
        'final_df': final_df,
        'fig_report': fig_report,
        'fig_report_buf': create_figure_buffer(fig_report),
        'fig_cal_buf': create_calendar_buffer(final_df),
        'input_filename': os.path.splitext(uploaded_file.name)[0],
        'k_final': k_final,
        'target_station': target_station,
    }


def store_analysis_result(result):
    previous_fig = st.session_state.get('fig_report')
    if previous_fig is not None:
        plt.close(previous_fig)

    for key, value in result.items():
        st.session_state[key] = value
    st.session_state['analyzed'] = True


def run_analysis(uploaded_file, floor_area, target_station, k_input):
    with st.spinner("解析レポートを作成中..."):
        result = build_analysis_result(uploaded_file, floor_area, target_station, k_input)
        if result is not None:
            store_analysis_result(result)


def render_hero():
    st.markdown(
        """
        <div class="hero-panel">
            <div class="hero-kicker">Power Clustering</div>
            <h1 class="hero-title">電力消費クラスタリングツール</h1>
            <p class="hero-copy">
                日別の消費パターンをクラスタリングし、気象条件や休日判定と合わせて可視化します。
                分析条件を左で指定すると、レポート生成とカレンダー出力まで一度に実行できます。
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_summary():
    final_df = st.session_state.final_df
    start_date = final_df.index.min().strftime('%Y-%m-%d')
    end_date = final_df.index.max().strftime('%Y-%m-%d')
    summary_items = [
        ("解析日数", f"{len(final_df)} days"),
        ("クラスタ数", str(st.session_state.k_final)),
        ("対象地点", st.session_state.target_station),
        ("期間", f"{start_date} to {end_date}"),
    ]
    cards = "".join(
        f'<div class="status-card"><div class="status-label">{label}</div><div class="status-value">{value}</div></div>'
        for label, value in summary_items
    )
    st.markdown(f'<div class="status-grid">{cards}</div>', unsafe_allow_html=True)


def render_sidebar_controls():
    with st.sidebar:
        st.markdown("## Analysis Setup")
        st.caption("CSV と分析条件を指定してレポートを生成します。")

        uploaded_file = st.file_uploader("電力データ CSV", type="csv")
        floor_area = st.number_input("延床面積 [m2]", value=5846, min_value=1)
        target_station = st.text_input("地点名", value="神戸")

        st.markdown("### Clustering")
        auto_k = st.checkbox("クラスタ数を自動決定", value=True)
        k_input = None if auto_k else st.slider("クラスタ数 (k)", 2, 10, 4)

        run_btn = st.button("分析を実行", use_container_width=True)

    return uploaded_file, floor_area, target_station, k_input, run_btn


def render_results():
    render_summary()

    st.sidebar.markdown("---")
    st.sidebar.download_button(
        label="Word レポートを保存",
        data=create_word_report_buffer(
            st.session_state.input_filename,
            st.session_state.target_station,
            st.session_state.k_final,
            st.session_state.fig_report_buf,
            st.session_state.fig_cal_buf,
        ),
        file_name=f"report_{st.session_state.input_filename}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        use_container_width=True,
    )

    report_tab, calendar_tab, data_tab = st.tabs(["Analysis Report", "Calendar Preview", "Merged Raw Data"])

    with report_tab:
        st.subheader("パターン分析レポート")
        st.pyplot(st.session_state.fig_report)

    with calendar_tab:
        st.subheader("カレンダープレビュー")
        st.image(st.session_state.fig_cal_buf, use_container_width=True)

    with data_tab:
        st.subheader("結合済み RAW データ")
        st.dataframe(st.session_state.final_df, use_container_width=True, height=460)


def render_empty_state():
    st.markdown(
        """
        <div class="empty-state">
            <div class="hero-kicker">Ready</div>
            <h3 style="margin:0;color:#1a1a1a;font-weight:600;">分析条件を設定するとここに結果が表示されます</h3>
            <p style="margin:1rem 0 0;color:#555550;line-height:1.85;font-size:0.93rem;">
                左のサイドバーで電力データ CSV、延床面積、地点名、クラスタ数の条件を指定して
                分析を実行してください。結果はレポート、カレンダー、結合データの3つのビューで確認できます。
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )


def main():
    inject_styles()
    ensure_session_state()
    if not check_password():
        st.stop()

    render_hero()
    uploaded_file, floor_area, target_station, k_input, run_btn = render_sidebar_controls()

    if run_btn:
        if uploaded_file is None:
            st.warning("CSV ファイルをアップロードしてください")
        else:
            try:
                run_analysis(uploaded_file, floor_area, target_station, k_input)
            except Exception as exc:
                st.error(f"エラーが発生しました: {exc}")
                st.exception(exc)

    if st.session_state.analyzed:
        render_results()
    else:
        render_empty_state()


main()